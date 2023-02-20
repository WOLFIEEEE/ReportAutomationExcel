from openpyxl.chart import BarChart, Reference, Series

def create_column_chart(sheet_name, chart_title, workbook):
    # get the worksheet by name
    sheet = workbook[sheet_name]

    # create a new chart
    chart = BarChart()

    # set the title of the chart
    chart.title = chart_title

    # define the data range for the chart
    data = Reference(sheet, min_col=2, min_row=2, max_row=sheet.max_row, max_col=2)

    # define the x-axis categories for the chart
    categories = Reference(sheet, min_col=1, min_row=2, max_row=sheet.max_row)

    # add the data and categories to the chart
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)

    # set the style of the chart
    chart.style = 10

    # add the chart to the worksheet
    sheet.add_chart(chart, "D1")


    return workbook
